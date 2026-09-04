#!/usr/bin/env python3
"""Render validated scientific Markdown as a styled, static-citation DOCX."""

from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse
import zipfile
from dataclasses import replace
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from docx.text.run import Run

from docx_common import (
    BLACK,
    BLUE,
    EMU_PER_INCH,
    GRAY,
    PROFILES,
    Profile,
    clear_document_body,
    configure_document,
    document_title,
    heading_size,
    horizontal_rule,
    linearize_math,
    paragraph_left_border,
    parse_link_target,
    prevent_row_split,
    repeat_table_header,
    set_cell_margins,
    set_cell_width,
    set_header_bottom_border,
    set_run_font,
    set_three_line_borders,
    shade_paragraph,
)
from md_common import (
    FENCE_RE,
    is_table_separator,
    plain_text,
    split_table_row,
    strip_front_matter,
)
from validate_markdown import validate


MERMAID_PACKAGE = "@mermaid-js/mermaid-cli@11.17.0"
HEADING_RE = re.compile(r"^\s*(#{1,6})\s+(.+?)\s*$")
LIST_RE = re.compile(
    r"^(?P<indent>\s*)(?:(?P<bullet>[-+*])|(?P<number>\d+)[.)])\s+(?P<text>.+)$"
)
IMAGE_RE = re.compile(r"^\s*!\[(?P<alt>[^]]*)\]\((?P<target>[^)]+)\)\s*$")
LINK_RE = re.compile(r"\[(?P<label>[^]\n]+)\]\((?P<target>[^)]+)\)")
INLINE_RE = re.compile(
    r"(?P<link>\[[^]\n]+\]\([^)]+\))"
    r"|(?P<bold>\*\*[^*\n]+\*\*)"
    r"|(?P<code>`[^`\n]+`)"
    r"|(?P<math>\$[^$\n]+\$)"
    r"|(?P<citation>(?<![\w!])\[\d+(?:\s*(?:,|[-–—])\s*\d+)*\](?!\s*\())"
    r"|(?P<italic>(?<!\*)\*[^*\n]+\*(?!\*))"
)


class Renderer:
    def __init__(
        self,
        document,
        source: Path,
        profile: Profile,
        *,
        mermaid_mode: str,
    ) -> None:
        self.document = document
        self.source = source
        self.profile = profile
        self.mermaid_mode = mermaid_mode
        self.warnings: list[str] = []
        self.table_count = 0
        self.figure_count = 0
        self.diagram_count = 0
        self.code_block_count = 0
        self.title_seen = False

    @property
    def usable_width(self) -> float:
        section = self.document.sections[-1]
        return (
            section.page_width - section.left_margin - section.right_margin
        ) / EMU_PER_INCH

    def add_text(
        self,
        paragraph,
        text: str,
        *,
        bold: bool = False,
        italic: bool = False,
        font: str | None = None,
        size: float | None = None,
        color: RGBColor = BLACK,
    ) -> Run:
        run = paragraph.add_run(text)
        set_run_font(
            run,
            font or self.profile.font,
            size or self.profile.font_size,
            bold=bold,
            italic=italic,
            color=color,
        )
        return run

    def add_hyperlink(self, paragraph, label: str, target: str, *, bold: bool) -> None:
        if not re.match(r"^(?:https?|mailto):", target, re.I):
            self.add_text(paragraph, label, bold=bold)
            return
        relationship = paragraph.part.relate_to(
            target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        )
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship)
        run_element = OxmlElement("w:r")
        hyperlink.append(run_element)
        paragraph._p.append(hyperlink)
        run = Run(run_element, paragraph)
        run.text = label
        set_run_font(
            run,
            self.profile.font,
            self.profile.font_size,
            bold=bold,
            color=BLUE,
            underline=True,
        )

    def add_math(self, paragraph, expression: str, *, bold: bool = False) -> None:
        value = linearize_math(expression)
        cursor = 0
        for match in re.finditer(r"([_^])(?:\{([^{}]+)\}|([^\s]))", value):
            if match.start() > cursor:
                self.add_text(
                    paragraph,
                    value[cursor : match.start()],
                    bold=bold,
                    font="Cambria Math",
                )
            token = match.group(2) or match.group(3) or ""
            run = self.add_text(
                paragraph,
                token,
                bold=bold,
                font="Cambria Math",
                size=self.profile.font_size * 0.88,
            )
            run.font.superscript = match.group(1) == "^"
            run.font.subscript = match.group(1) == "_"
            cursor = match.end()
        if cursor < len(value):
            self.add_text(paragraph, value[cursor:], bold=bold, font="Cambria Math")

    def add_inline(
        self,
        paragraph,
        text: str,
        *,
        base_bold: bool = False,
        size: float | None = None,
    ) -> None:
        text = text.replace(r"\($", "$").replace(r"$\)", "$")
        cursor = 0
        for match in INLINE_RE.finditer(text):
            if match.start() > cursor:
                self.add_plain_with_breaks(
                    paragraph, text[cursor : match.start()], bold=base_bold, size=size
                )
            token = match.group(0)
            if match.lastgroup == "link":
                link = LINK_RE.fullmatch(token)
                if link:
                    self.add_hyperlink(
                        paragraph,
                        link.group("label"),
                        parse_link_target(link.group("target")),
                        bold=base_bold,
                    )
            elif match.lastgroup == "bold":
                self.add_text(paragraph, token[2:-2], bold=True, size=size)
            elif match.lastgroup == "italic":
                self.add_text(paragraph, token[1:-1], bold=base_bold, italic=True, size=size)
            elif match.lastgroup == "code":
                run = self.add_text(
                    paragraph,
                    token[1:-1],
                    bold=base_bold,
                    font="Courier New",
                    size=(size or self.profile.font_size) * 0.92,
                )
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EFEFEF")
                run._element.get_or_add_rPr().append(shading)
            elif match.lastgroup == "math":
                self.add_math(paragraph, token[1:-1], bold=base_bold)
            elif match.lastgroup == "citation":
                if self.profile.citation_style == "superscript":
                    run = self.add_text(
                        paragraph,
                        token[1:-1],
                        size=(size or self.profile.font_size) * 0.82,
                    )
                    run.font.superscript = True
                else:
                    self.add_text(paragraph, token, bold=base_bold, size=size)
            cursor = match.end()
        if cursor < len(text):
            self.add_plain_with_breaks(paragraph, text[cursor:], bold=base_bold, size=size)

    def add_plain_with_breaks(
        self, paragraph, text: str, *, bold: bool = False, size: float | None = None
    ) -> None:
        parts = re.split(r"<br\s*/?>", text, flags=re.I)
        for index, part in enumerate(parts):
            if part:
                self.add_text(paragraph, part, bold=bold, size=size)
            if index + 1 < len(parts):
                paragraph.add_run().add_break()

    def add_heading(self, level: int, text: str) -> None:
        is_title = level == 1 and not self.title_seen
        if is_title:
            style = "Title"
            self.title_seen = True
        else:
            style = f"Heading {min(level, 6)}"
        paragraph = self.document.add_paragraph(style=style)
        self.add_inline(
            paragraph,
            text,
            base_bold=True,
            size=heading_size(self.profile, level, is_title),
        )
        paragraph.paragraph_format.keep_with_next = True

    def add_paragraph(self, text: str) -> None:
        paragraph = self.document.add_paragraph()
        self.add_inline(paragraph, text)

    def add_list_item(self, match: re.Match[str]) -> None:
        level = min(3, len(match.group("indent").replace("\t", "    ")) // 2)
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28 + 0.20 * level)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(2)
        marker = "•" if match.group("bullet") else f"{match.group('number')}."
        self.add_text(paragraph, marker + " ", bold=False)
        self.add_inline(paragraph, match.group("text"))

    def add_blockquote(self, lines: Sequence[str]) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph_left_border(paragraph)
        value = " ".join(re.sub(r"^\s*>\s?", "", line).strip() for line in lines)
        self.add_inline(paragraph, value)

    def column_widths(self, rows: Sequence[Sequence[str]]) -> list[float]:
        columns = max(len(row) for row in rows)
        measures: list[float] = []
        for column in range(columns):
            maximum = max(
                len(plain_text(row[column])) if column < len(row) else 0
                for row in rows
            )
            measures.append(max(7.0, min(65.0, maximum)) ** 0.72)
        total = sum(measures)
        widths = [self.usable_width * value / total for value in measures]
        floor = min(0.58, self.usable_width / columns * 0.72)
        widths = [max(floor, width) for width in widths]
        scale = self.usable_width / sum(widths)
        return [width * scale for width in widths]

    def add_table(self, header: list[str], rows: list[list[str]], separator: str) -> None:
        values = [header, *rows]
        columns = len(header)
        table = self.document.add_table(rows=len(values), cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        if self.profile.table_style == "grid":
            table.style = "Table Grid"
        else:
            table.style = None
            set_three_line_borders(table)
        repeat_table_header(table.rows[0])
        widths = self.column_widths(values)

        alignments: list[WD_ALIGN_PARAGRAPH] = []
        for cell in split_table_row(separator):
            left, right = cell.startswith(":"), cell.endswith(":")
            if left and right:
                alignments.append(WD_ALIGN_PARAGRAPH.CENTER)
            elif right:
                alignments.append(WD_ALIGN_PARAGRAPH.RIGHT)
            else:
                alignments.append(WD_ALIGN_PARAGRAPH.LEFT)

        cell_size = self.profile.font_size
        if columns >= 5:
            cell_size = max(8.0, cell_size - 2.0)
            self.warnings.append(
                f"table with {columns} columns used {cell_size:g}-point text"
            )
        elif columns == 4:
            cell_size = max(8.5, cell_size - 1.0)

        for row_index, source_row in enumerate(values):
            row = table.rows[row_index]
            prevent_row_split(row)
            for column_index, cell in enumerate(row.cells):
                set_cell_width(cell, widths[column_index])
                set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if row_index == 0 and self.profile.table_style == "three-line":
                    set_header_bottom_border(cell)
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = alignments[column_index]
                self.add_inline(
                    paragraph,
                    source_row[column_index],
                    base_bold=row_index == 0,
                    size=cell_size,
                )
        self.table_count += 1
        spacer = self.document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)

    def add_code_block(self, language: str, lines: Sequence[str]) -> None:
        if language:
            label = self.document.add_paragraph()
            label.paragraph_format.space_before = Pt(3)
            label.paragraph_format.space_after = Pt(1)
            self.add_text(
                label,
                language,
                italic=True,
                font="Courier New",
                size=max(8.0, self.profile.font_size - 2.0),
                color=GRAY,
            )
        for index, line in enumerate(lines or [""]):
            paragraph = self.document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.16)
            paragraph.paragraph_format.right_indent = Inches(0.08)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0 if index + 1 < len(lines) else 5)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_together = True
            shade_paragraph(paragraph, "F2F2F2")
            self.add_text(
                paragraph,
                line if line else " ",
                font="Courier New",
                size=max(8.0, self.profile.font_size - 1.5),
            )
        self.code_block_count += 1

    def mermaid_command(self) -> list[str]:
        direct = shutil.which("mmdc")
        if direct:
            return [direct]
        npx = shutil.which("npx")
        if npx:
            return [npx, "--yes", MERMAID_PACKAGE]
        raise RuntimeError("Mermaid rendering requires mmdc or npx")

    def add_mermaid(self, lines: Sequence[str]) -> None:
        if self.mermaid_mode == "code":
            self.add_code_block("mermaid", lines)
            return
        try:
            with tempfile.TemporaryDirectory(prefix="scientific-md-mermaid-") as folder:
                source_path = Path(folder) / "workflow.mmd"
                output_path = Path(folder) / "workflow.png"
                source_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
                environment = os.environ.copy()
                if not environment.get("PUPPETEER_EXECUTABLE_PATH"):
                    for browser_name in ("google-chrome", "chromium", "chromium-browser"):
                        browser = shutil.which(browser_name)
                        if browser:
                            environment["PUPPETEER_EXECUTABLE_PATH"] = browser
                            break
                result = subprocess.run(
                    [
                        *self.mermaid_command(),
                        "--input",
                        str(source_path),
                        "--output",
                        str(output_path),
                        "--backgroundColor",
                        "white",
                        "--scale",
                        "2",
                        "--quiet",
                    ],
                    text=True,
                    capture_output=True,
                    env=environment,
                    timeout=120,
                    check=False,
                )
                if result.returncode or not output_path.is_file():
                    detail = (result.stderr or result.stdout).strip()
                    raise RuntimeError(detail or "Mermaid produced no image")
                if output_path.read_bytes()[:8] != b"\x89PNG\r\n\x1a\n":
                    raise RuntimeError("Mermaid output is not a valid PNG")
                paragraph = self.document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run()
                run.add_picture(str(output_path), width=Inches(self.usable_width))
                self.figure_count += 1
                self.diagram_count += 1
        except (OSError, RuntimeError, subprocess.SubprocessError) as error:
            if self.mermaid_mode == "render":
                raise RuntimeError(f"Mermaid rendering failed: {error}") from error
            self.warnings.append(
                f"Mermaid rendering failed; source preserved as code: {error}"
            )
            self.add_code_block("mermaid", lines)

    def add_image(self, alt: str, raw_target: str) -> None:
        target = urllib.parse.unquote(parse_link_target(raw_target))
        if re.match(r"^https?://", target, re.I):
            self.warnings.append(f"remote image not embedded: {target}")
            self.add_paragraph(f"Image not embedded: {alt or target} ({target})")
            return
        path = (self.source.parent / target).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Image does not exist: {target}")
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(path), width=Inches(min(6.5, self.usable_width)))
        self.figure_count += 1
        if alt:
            caption = self.document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.add_text(
                caption,
                alt,
                italic=True,
                size=max(8.5, self.profile.font_size - 1.0),
                color=GRAY,
            )

    def is_block_start(self, lines: Sequence[str], index: int) -> bool:
        stripped = lines[index].strip()
        if not stripped:
            return True
        if (
            FENCE_RE.match(lines[index])
            or HEADING_RE.match(lines[index])
            or LIST_RE.match(lines[index])
            or IMAGE_RE.match(lines[index])
            or stripped.startswith(">")
            or stripped in {"---", "***", "___"}
            or stripped.startswith("<!--")
        ):
            return True
        return (
            "|" in lines[index]
            and index + 1 < len(lines)
            and is_table_separator(lines[index + 1])
        )

    def render(self, source_lines: Sequence[str]) -> None:
        lines, _offset = strip_front_matter(source_lines)
        index = 0
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            if not stripped:
                index += 1
                continue

            if stripped.startswith("<!--"):
                while index < len(lines):
                    if "-->" in lines[index]:
                        index += 1
                        break
                    index += 1
                continue

            fence_match = FENCE_RE.match(line)
            if fence_match:
                fence = fence_match.group("fence")
                language = fence_match.group("info").strip().split(maxsplit=1)[0]
                index += 1
                block: list[str] = []
                close_re = re.compile(rf"^\s*{re.escape(fence[0])}{{{len(fence)},}}\s*$")
                while index < len(lines) and not close_re.match(lines[index]):
                    block.append(lines[index])
                    index += 1
                if index >= len(lines):
                    raise ValueError("Unclosed fenced block")
                if language.lower() == "mermaid":
                    self.add_mermaid(block)
                else:
                    self.add_code_block(language, block)
                index += 1
                continue

            heading = HEADING_RE.match(line)
            if heading:
                self.add_heading(len(heading.group(1)), heading.group(2))
                index += 1
                continue

            if stripped in {"---", "***", "___"}:
                paragraph = self.document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(5)
                horizontal_rule(paragraph)
                index += 1
                continue

            image = IMAGE_RE.match(line)
            if image:
                self.add_image(image.group("alt"), image.group("target"))
                index += 1
                continue

            if (
                "|" in line
                and index + 1 < len(lines)
                and is_table_separator(lines[index + 1])
            ):
                header = split_table_row(line)
                separator = lines[index + 1]
                index += 2
                rows: list[list[str]] = []
                while index < len(lines) and lines[index].strip() and "|" in lines[index]:
                    row = split_table_row(lines[index])
                    if len(row) != len(header):
                        raise ValueError(
                            f"Table row at Markdown line {index + 1} has "
                            f"{len(row)} cells; expected {len(header)}"
                        )
                    rows.append(row)
                    index += 1
                self.add_table(header, rows, separator)
                continue

            if stripped.startswith(">"):
                quote: list[str] = []
                while index < len(lines) and lines[index].strip().startswith(">"):
                    quote.append(lines[index])
                    index += 1
                self.add_blockquote(quote)
                continue

            list_match = LIST_RE.match(line)
            if list_match:
                self.add_list_item(list_match)
                index += 1
                continue

            paragraph_lines: list[str] = []
            while index < len(lines) and not self.is_block_start(lines, index):
                raw = lines[index].rstrip("\r\n")
                hard_break = len(raw) - len(raw.rstrip(" ")) >= 2
                value = raw.strip()
                if hard_break:
                    value += "<br>"
                paragraph_lines.append(value)
                index += 1
            if not paragraph_lines:
                self.add_paragraph(stripped)
                index += 1
                continue
            text = " ".join(paragraph_lines).replace("<br> ", "<br>")
            self.add_paragraph(text)


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Markdown source")
    parser.add_argument("-o", "--output", type=Path, help="DOCX destination")
    parser.add_argument("--profile", choices=sorted(PROFILES), default="nature")
    parser.add_argument("--font")
    parser.add_argument("--font-size", type=float)
    parser.add_argument("--citation-style", choices=("brackets", "superscript"))
    parser.add_argument("--table-style", choices=("three-line", "grid"))
    parser.add_argument("--template", type=Path, help="Clean DOCX style template")
    parser.add_argument("--landscape", action="store_true")
    parser.add_argument("--mermaid", choices=("auto", "render", "code"), default="auto")
    parser.add_argument("--title", help="Override Word metadata title")
    parser.add_argument("--no-page-numbers", action="store_true")
    parser.add_argument("--strict-markdown", action="store_true")
    parser.add_argument("--force", action="store_true", help="Replace an existing output")
    return parser.parse_args()


def main() -> int:
    args = arguments()
    source = args.input.resolve()
    if not source.is_file():
        print(f"error: input does not exist: {source}", file=sys.stderr)
        return 2
    output = (args.output or source.with_suffix(".docx")).resolve()
    if output == source:
        print("error: output must differ from the Markdown input", file=sys.stderr)
        return 2
    if output.exists() and not args.force:
        print(f"error: output already exists (use --force): {output}", file=sys.stderr)
        return 2
    if args.template and not args.template.is_file():
        print(f"error: template does not exist: {args.template}", file=sys.stderr)
        return 2
    if args.font_size is not None and not 7 <= args.font_size <= 24:
        print("error: --font-size must be between 7 and 24 points", file=sys.stderr)
        return 2

    issues = validate(source)
    errors = [issue for issue in issues if issue.severity == "error"]
    warnings = [issue for issue in issues if issue.severity == "warning"]
    if errors or (args.strict_markdown and warnings):
        for issue in issues:
            location = f":{issue.line}" if issue.line is not None else ""
            print(
                f"{source}{location}: {issue.severity}: {issue.code}: {issue.message}",
                file=sys.stderr,
            )
        return 2
    for issue in warnings:
        location = f":{issue.line}" if issue.line is not None else ""
        print(
            f"warning: {source}{location}: {issue.code}: {issue.message}",
            file=sys.stderr,
        )

    profile = PROFILES[args.profile]
    profile = replace(
        profile,
        font=args.font or profile.font,
        font_size=args.font_size or profile.font_size,
        citation_style=args.citation_style or profile.citation_style,
        table_style=args.table_style or profile.table_style,
    )

    try:
        text = source.read_text(encoding="utf-8")
        lines = text.splitlines()
        if args.template:
            document = Document(args.template)
            clear_document_body(document)
        else:
            document = Document()
        configure_document(
            document,
            profile,
            landscape=args.landscape,
            page_numbers=not args.no_page_numbers,
        )
        title = args.title or document_title(lines, source.stem.replace("_", " "))
        document.core_properties.title = title
        document.core_properties.subject = "Evidence-grounded scientific document"
        document.core_properties.keywords = "scientific writing, evidence, Markdown"
        document.core_properties.comments = (
            f"Generated from {source.name} by the scientific-md-docx skill; "
            "citations are static."
        )

        renderer = Renderer(
            document,
            source,
            profile,
            mermaid_mode=args.mermaid,
        )
        renderer.render(lines)

        output.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(
            prefix=f".{output.stem}.", suffix=".docx", dir=output.parent
        )
        os.close(descriptor)
        staging = Path(temporary_name)
        try:
            document.save(staging)
            with zipfile.ZipFile(staging, "r") as archive:
                bad_member = archive.testzip()
                if bad_member:
                    raise ValueError(f"Corrupt DOCX member: {bad_member}")
            reopened = Document(staging)
            if len(reopened.tables) != renderer.table_count:
                raise ValueError(
                    "Generated table count changed after reopening the DOCX"
                )
            if len(reopened.inline_shapes) != renderer.figure_count:
                raise ValueError(
                    "Generated figure count changed after reopening the DOCX"
                )
            os.replace(staging, output)
            os.chmod(output, 0o664)
        finally:
            if staging.exists():
                staging.unlink()
    except (OSError, ValueError, RuntimeError, zipfile.BadZipFile) as error:
        print(f"error: {error}", file=sys.stderr)
        return 2

    for warning in renderer.warnings:
        print(f"warning: {warning}", file=sys.stderr)
    print(
        f"created {output}: {renderer.table_count} tables, "
        f"{renderer.figure_count} figures ({renderer.diagram_count} Mermaid), "
        f"{renderer.code_block_count} code blocks"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
